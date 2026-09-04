"""Tests for documents/docx_comments.py — the stopgap that restores
comments Casual Docs' reimport silently drops (see that module's
docstring for the full diagnosis). Deliberately its own test module
(rather than folded into documents/tests.py) given how much fixture
machinery a faithful test needs — python-docx's own add_comment() can't
build the modern 4-part comment schema this targets at all, so the
fixtures here are hand-built directly from the real structure confirmed
against a genuine saved Casual Docs file.
"""
import io
import os
import tempfile
import zipfile

from django.test import TestCase
from docx import Document as DocxDocument

from documents.docx_comments import restore_missing_comments, strip_comments_by_author

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
_W15 = "http://schemas.microsoft.com/office/word/2012/wordml"
_W16CID = "http://schemas.microsoft.com/office/word/2016/wordml/cid"
_W16CEX = "http://schemas.microsoft.com/office/word/2018/wordml/cex"


def _build_docx(anchor_ids, comments=None):
    """A minimal but structurally real .docx.

    anchor_ids: comment IDs to anchor in document.xml (one paragraph per
        id, with commentRangeStart/commentRangeEnd/commentReference) —
        this is what survives Casual Docs' reimport (confirmed live: the
        highlight shows even when the comment panel doesn't).
    comments: optional {id: text} — if given, builds the full real 4-part
        comment schema (comments.xml + commentsExtended.xml +
        commentsIds.xml + commentsExtensible.xml, correctly
        cross-referenced by paraId/durableId) for exactly those ids. Omit
        entirely (None/empty) to simulate the actual bug: anchors present,
        comment data entirely absent.
    """
    comments = comments or {}
    body_paras = []
    for cid in anchor_ids:
        body_paras.append(
            f'<w:p><w:commentRangeStart w:id="{cid}"/>'
            f'<w:r><w:t>Paragraph anchored to comment {cid}.</w:t></w:r>'
            f'<w:commentRangeEnd w:id="{cid}"/>'
            f'<w:r><w:commentReference w:id="{cid}"/></w:r></w:p>'
        )
    document_xml = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<w:document xmlns:w="{_W}"><w:body>{"".join(body_paras)}'
        f'<w:sectPr/></w:body></w:document>'
    ).encode("utf-8")

    content_types = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    )

    parts = {
        "_rels/.rels": (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            "</Relationships>"
        ).encode("utf-8"),
        "word/document.xml": document_xml,
    }

    if comments:
        content_types += (
            '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
            '<Override PartName="/word/commentsExtended.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtended+xml"/>'
            '<Override PartName="/word/commentsIds.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.commentsIds+xml"/>'
            '<Override PartName="/word/commentsExtensible.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.commentsExtensible+xml"/>'
        )
        doc_rels += (
            '<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
            '<Relationship Id="rId6" Type="http://schemas.microsoft.com/office/2011/relationships/commentsExtended" Target="commentsExtended.xml"/>'
            '<Relationship Id="rId7" Type="http://schemas.microsoft.com/office/2016/09/relationships/commentsIds" Target="commentsIds.xml"/>'
            '<Relationship Id="rId8" Type="http://schemas.microsoft.com/office/2018/08/relationships/commentsExtensible" Target="commentsExtensible.xml"/>'
        )

        comment_els, ex_els, ids_els, cex_els = [], [], [], []
        for i, (cid, value) in enumerate(comments.items()):
            # value is either plain text (default author "Test") or a
            # (text, author) tuple, for tests that need a specific author.
            text, author = value if isinstance(value, tuple) else (value, "Test")
            para_id = f"{i:08X}"
            durable_id = f"{i + 1:08X}"
            comment_els.append(
                f'<w:comment w:id="{cid}" w:author="{author}" w:date="2026-01-01T00:00:00Z">'
                f'<w:p w14:paraId="{para_id}"><w:r><w:t>{text}</w:t></w:r></w:p></w:comment>'
            )
            ex_els.append(f'<w15:commentEx w15:paraId="{para_id}" w15:done="0"/>')
            ids_els.append(f'<w16cid:commentId w16cid:paraId="{para_id}" w16cid:durableId="{durable_id}"/>')
            cex_els.append(f'<w16cex:commentExtensible w16cex:durableId="{durable_id}" w16cex:dateUtc="2026-01-01T00:00:00Z"/>')

        parts["word/comments.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w:comments xmlns:w="{_W}" xmlns:w14="{_W14}">{"".join(comment_els)}</w:comments>'
        ).encode("utf-8")
        parts["word/commentsExtended.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w15:commentsEx xmlns:w15="{_W15}">{"".join(ex_els)}</w15:commentsEx>'
        ).encode("utf-8")
        parts["word/commentsIds.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w16cid:commentsIds xmlns:w16cid="{_W16CID}">{"".join(ids_els)}</w16cid:commentsIds>'
        ).encode("utf-8")
        parts["word/commentsExtensible.xml"] = (
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            f'<w16cex:commentsExtensible xmlns:w16cex="{_W16CEX}">{"".join(cex_els)}</w16cex:commentsExtensible>'
        ).encode("utf-8")

    parts["[Content_Types].xml"] = (content_types + "</Types>").encode("utf-8")
    parts["word/_rels/document.xml.rels"] = (doc_rels + "</Relationships>").encode("utf-8")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts.items():
            z.writestr(name, data)
    return buf.getvalue()


def _comment_texts(docx_bytes):
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        try:
            comments_xml = z.read("word/comments.xml")
        except KeyError:
            return {}
    from lxml import etree
    root = etree.fromstring(comments_xml)
    out = {}
    for el in root.findall(f"{{{_W}}}comment"):
        text_el = el.find(f".//{{{_W}}}t")
        out[el.get(f"{{{_W}}}id")] = text_el.text if text_el is not None else None
    return out


class RestoreMissingCommentsTests(TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="lepmits_test_comments_")
        self._existing_path = os.path.join(self._tmp, "existing.docx")

    def _write_existing(self, docx_bytes):
        with open(self._existing_path, "wb") as f:
            f.write(docx_bytes)

    def test_no_existing_file_is_a_noop(self):
        new_bytes = _build_docx(anchor_ids=["19"])
        out = restore_missing_comments(
            os.path.join(self._tmp, "does-not-exist.docx"), new_bytes
        )
        self.assertEqual(out, new_bytes)

    def test_existing_file_without_comments_is_a_noop(self):
        self._write_existing(_build_docx(anchor_ids=[]))
        new_bytes = _build_docx(anchor_ids=["19"])
        out = restore_missing_comments(self._existing_path, new_bytes)
        self.assertEqual(out, new_bytes)

    def test_restores_comment_lost_on_reimport(self):
        # The actual reported bug: anchor survives, comment body doesn't.
        self._write_existing(_build_docx(anchor_ids=["19"], comments={"19": "Original comment text."}))
        new_bytes = _build_docx(anchor_ids=["19"])  # anchor present, no comment parts at all

        out = restore_missing_comments(self._existing_path, new_bytes)

        self.assertNotEqual(out, new_bytes)
        texts = _comment_texts(out)
        self.assertEqual(texts.get("19"), "Original comment text.")
        # Result must still be a valid, openable docx.
        DocxDocument(io.BytesIO(out))

    def test_does_not_restore_when_ids_dont_overlap(self):
        self._write_existing(_build_docx(anchor_ids=["19"], comments={"19": "Original."}))
        new_bytes = _build_docx(anchor_ids=["99"])  # different id, no overlap

        out = restore_missing_comments(self._existing_path, new_bytes)
        self.assertEqual(out, new_bytes)

    def test_does_not_duplicate_when_new_file_already_has_it(self):
        self._write_existing(_build_docx(anchor_ids=["19"], comments={"19": "Original."}))
        new_bytes = _build_docx(anchor_ids=["19"], comments={"19": "Already present, edited."})

        out = restore_missing_comments(self._existing_path, new_bytes)
        self.assertEqual(out, new_bytes)  # already has id 19 -> nothing to restore
        self.assertEqual(_comment_texts(out).get("19"), "Already present, edited.")

    def test_restores_only_the_missing_one_of_several(self):
        self._write_existing(_build_docx(
            anchor_ids=["19", "20"],
            comments={"19": "First comment.", "20": "Second comment."},
        ))
        # New save kept both anchors but only id 19's comment data survived.
        new_bytes = _build_docx(anchor_ids=["19", "20"], comments={"19": "First comment."})

        out = restore_missing_comments(self._existing_path, new_bytes)

        texts = _comment_texts(out)
        self.assertEqual(texts.get("19"), "First comment.")
        self.assertEqual(texts.get("20"), "Second comment.")
        DocxDocument(io.BytesIO(out))

    def test_result_preserves_unrelated_parts(self):
        self._write_existing(_build_docx(anchor_ids=["19"], comments={"19": "Original."}))
        new_bytes = _build_docx(anchor_ids=["19"])

        out = restore_missing_comments(self._existing_path, new_bytes)

        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        self.assertIn('Paragraph anchored to comment 19.', doc_xml)


class StripCommentsByAuthorTests(TestCase):
    def test_noop_when_nothing_anchored_or_commented(self):
        docx = _build_docx(anchor_ids=[])
        out, count = strip_comments_by_author(docx)
        self.assertEqual(out, docx)
        self.assertEqual(count, 0)

    def test_strips_orphaned_anchor_with_no_comments_part_at_all(self):
        # The Casual Docs reimport-drops-comment-body bug (see
        # restore_missing_comments' docstring), caught a save too late for
        # that stopgap to have a same-session chance to restore it:
        # document.xml still anchors (and highlights) the range, but
        # comments.xml never made it back in — not even an empty one, the
        # whole part is gone. Author unknowable since there's nothing left
        # to check it against; safe to strip anyway per this module's
        # comment above AI_COMMENT_AUTHORS (no human "add comment" UI
        # exists in this app, so an orphaned anchor can never be a real
        # reviewer's comment).
        docx = _build_docx(anchor_ids=["19"])
        out, count = strip_comments_by_author(docx)

        self.assertEqual(count, 1)
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        self.assertNotIn("commentRangeStart", doc_xml)
        self.assertNotIn("commentRangeEnd", doc_xml)
        self.assertNotIn("commentReference", doc_xml)
        # The paragraph text itself is untouched — only the comment anchor.
        self.assertIn("Paragraph anchored to comment 19.", doc_xml)
        DocxDocument(io.BytesIO(out))

    def test_orphaned_anchor_stripped_but_real_other_comment_kept(self):
        # Mixed case: id 19 has no comment body anywhere (orphaned, gets
        # stripped); id 20 has a real body from a non-AI author (kept) —
        # confirms orphan-detection doesn't get confused with author
        # matching and doesn't over-strip a comment that's actually there.
        docx = _build_docx(
            anchor_ids=["19", "20"],
            comments={"20": ("Human note.", "A Reviewer")},
        )
        out, count = strip_comments_by_author(docx)

        self.assertEqual(count, 1)
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        self.assertNotIn('w:id="19"', doc_xml)
        self.assertIn('w:id="20"', doc_xml)
        self.assertEqual(_comment_texts(out).get("20"), "Human note.")

    def test_noop_when_no_matching_author(self):
        docx = _build_docx(anchor_ids=["19"], comments={"19": ("Human note.", "A Reviewer")})
        out, count = strip_comments_by_author(docx)
        self.assertEqual(out, docx)
        self.assertEqual(count, 0)

    def test_strips_ai_comment_and_its_anchors(self):
        docx = _build_docx(
            anchor_ids=["19"],
            comments={"19": ("Possible similarity to X.", "Similarity Check (AI)")},
        )
        out, count = strip_comments_by_author(docx)

        self.assertEqual(count, 1)
        self.assertEqual(_comment_texts(out), {})
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        self.assertNotIn("commentRangeStart", doc_xml)
        self.assertNotIn("commentRangeEnd", doc_xml)
        self.assertNotIn("commentReference", doc_xml)
        # The paragraph text itself is untouched — only the comment.
        self.assertIn("Paragraph anchored to comment 19.", doc_xml)
        DocxDocument(io.BytesIO(out))

    def test_strips_server_side_author_too(self):
        docx = _build_docx(
            anchor_ids=["19"],
            comments={"19": ("Possible similarity to X.", "LePMITS AI")},
        )
        out, count = strip_comments_by_author(docx)
        self.assertEqual(count, 1)
        self.assertEqual(_comment_texts(out), {})

    def test_leaves_non_ai_comment_untouched(self):
        docx = _build_docx(
            anchor_ids=["19", "20"],
            comments={
                "19": ("AI flag.", "Similarity Check (AI)"),
                "20": ("Human note.", "A Reviewer"),
            },
        )
        out, count = strip_comments_by_author(docx)

        self.assertEqual(count, 1)
        texts = _comment_texts(out)
        self.assertNotIn("19", texts)
        self.assertEqual(texts.get("20"), "Human note.")
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            doc_xml = z.read("word/document.xml").decode("utf-8")
        # Anchor 20's markers survive; anchor 19's are gone.
        self.assertIn('w:id="20"', doc_xml)
        self.assertNotIn('w:id="19"', doc_xml)
        DocxDocument(io.BytesIO(out))

    def test_removes_cross_referenced_extended_parts(self):
        docx = _build_docx(
            anchor_ids=["19"],
            comments={"19": ("AI flag.", "Similarity Check (AI)")},
        )
        out, count = strip_comments_by_author(docx)
        self.assertEqual(count, 1)
        with zipfile.ZipFile(io.BytesIO(out)) as z:
            ext_xml = z.read("word/commentsExtended.xml").decode("utf-8")
            ids_xml = z.read("word/commentsIds.xml").decode("utf-8")
            cex_xml = z.read("word/commentsExtensible.xml").decode("utf-8")
        self.assertNotIn("commentEx ", ext_xml)
        self.assertNotIn("commentId ", ids_xml)
        self.assertNotIn("commentExtensible ", cex_xml)
