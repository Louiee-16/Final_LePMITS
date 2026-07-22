from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

# ── Styles ────────────────────────────────────────────────────────────────────
normal_style = doc.styles['Normal']
normal_style.font.name = 'Calibri'
normal_style.font.size = Pt(10)

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def set_cell_border(cell):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'),   'single')
        border.set(qn('w:sz'),    '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'D1D5DB')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_table(doc, title, headers, rows):
    # Table title
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold      = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x7C, 0x0D, 0x0E)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    col_widths = [Cm(4), Cm(4.5), Cm(7), Cm(5.5)]

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_bg(hdr_cells[i], '1E293B')
        set_cell_border(hdr_cells[i])
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(h)
        run.bold            = True
        run.font.color.rgb  = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size       = Pt(9)
        p.alignment         = WD_ALIGN_PARAGRAPH.CENTER

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg = 'F8FAFC' if r_idx % 2 == 0 else 'FFFFFF'
        for c_idx, cell_text in enumerate(row_data):
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_bg(row_cells[c_idx], bg)
            set_cell_border(row_cells[c_idx])
            p = row_cells[c_idx].paragraphs[0]
            run = p.add_run(str(cell_text))
            run.font.size = Pt(9)
            if c_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph()

# ── Title page ────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('DATA DICTIONARY')
run.bold            = True
run.font.size       = Pt(18)
run.font.color.rgb  = RGBColor(0x7C, 0x0D, 0x0E)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('Legislative Document Management System\nCity of San Juan, Metro Manila')
r2.font.size      = Pt(11)
r2.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

doc.add_paragraph()

# ═══════════════════════════════════════════════════════════════════════════════
HEADERS = ['Field Name', 'Data Type', 'Description', 'Constraints']

# ── accounts_user ─────────────────────────────────────────────────────────────
add_table(doc, 'Table: accounts_user', HEADERS, [
    ['id',               'AutoField',    'Primary key',                             'Auto-generated'],
    ['username',         'CharField',    'Login username',                          'Max: 150, Unique, Required'],
    ['password',         'CharField',    'Hashed password',                         'Required'],
    ['email',            'CharField',    'Email address',                           'Optional'],
    ['first_name',       'CharField',    'First name',                              'Max: 150, Optional'],
    ['last_name',        'CharField',    'Last name',                               'Max: 150, Optional'],
    ['is_staff',         'BooleanField', 'Django admin access flag',                'Default: False'],
    ['is_active',        'BooleanField', 'Account active flag',                     'Default: True'],
    ['date_joined',      'DateTimeField','Account creation date',                   'Auto-generated'],
    ['role',             'CharField',    'User role in the system',                 'Max: 20, Choices: SECRETARIAT, STAFF, COUNCILOR, BARANGAY, ADMIN, Default: STAFF'],
    ['office_or_district','CharField',   'Office or district assignment',           'Max: 100, Optional'],
])

# ── councilors_councilor ──────────────────────────────────────────────────────
add_table(doc, 'Table: councilors_councilor', HEADERS, [
    ['id',              'AutoField',    'Primary key',                'Auto-generated'],
    ['user_id',         'ForeignKey (accounts_user)', 'Linked user account', 'OneToOne, Cascade delete'],
    ['name',            'CharField',    'Full name of councilor',     'Max: 50, Optional'],
    ['email',           'EmailField',   'Email address',              'Max: 50, Required'],
    ['district',        'IntegerField', 'District assignment',        'Choices: 1, 2, Required'],
    ['profile_picture', 'ImageField',   'Profile photo',              'Upload to: councilors/, Optional'],
    ['is_active',       'BooleanField', 'Currently in office flag',   'Default: True'],
])

# ── committees_committee ──────────────────────────────────────────────────────
add_table(doc, 'Table: committees_committee', HEADERS, [
    ['id',               'AutoField',    'Primary key',            'Auto-generated'],
    ['name',             'CharField',    'Name of the committee',  'Max: 255, Unique, Required'],
    ['description',      'TextField',    'Committee description',  'Optional'],
    ['chairman_id',      'ForeignKey (councilors_councilor)', 'Committee chairman',       'Set null on delete, Optional'],
    ['vice_chairman_id', 'ForeignKey (councilors_councilor)', 'Committee vice chairman',  'Set null on delete, Optional'],
    ['member_id',        'ForeignKey (councilors_councilor)', 'Committee member',         'Set null on delete, Optional'],
    ['created_at',       'DateTimeField','Date committee was created','Auto-generated'],
])

# ── barangay_barangay ─────────────────────────────────────────────────────────
add_table(doc, 'Table: barangay_barangay', HEADERS, [
    ['id',             'AutoField',    'Primary key',              'Auto-generated'],
    ['user_id',        'ForeignKey (accounts_user)', 'Linked user account', 'OneToOne, Cascade delete'],
    ['barangay_name',  'CharField',    'Name of the barangay',     'Max: 100, Unique, Required'],
    ['captain',        'CharField',    'Name of barangay captain', 'Max: 255, Required'],
    ['email',          'EmailField',   'Barangay email address',   'Required'],
    ['is_active',      'BooleanField', 'Active status flag',       'Default: True'],
])

# ── barangay_barangayfiles ────────────────────────────────────────────────────
add_table(doc, 'Table: barangay_barangayfiles', HEADERS, [
    ['id',                   'AutoField',    'Primary key',                       'Auto-generated'],
    ['origin_barangay_id',   'ForeignKey (barangay_barangay)', 'Source barangay', 'Cascade delete, Required'],
    ['scanned_pdf',          'FileField',    'Uploaded scanned document',         'Upload to: barangay_scans/YYYY/, Optional'],
    ['date_submitted',       'DateTimeField','Submission timestamp',              'Auto-generated'],
    ['remarks',              'TextField',    'Additional remarks',                'Max: 600, Optional'],
    ['title',                'CharField',    'Document title',                    'Max: 100, Optional'],
    ['subject',              'CharField',    'Document subject',                  'Max: 200, Optional'],
    ['status',               'CharField',    'Current processing status',         'Max: 50, Required'],
    ['referred_committee_id','ForeignKey (committees_committee)', 'Referred committee', 'Set null on delete, Optional'],
])

# ── archives_archives ─────────────────────────────────────────────────────────
add_table(doc, 'Table: archives_archives', HEADERS, [
    ['id',           'AutoField',    'Primary key',                              'Auto-generated'],
    ['original_doc_id','ForeignKey (documents_document)', 'Reference to original document', 'Cascade delete, Required'],
    ['title',        'CharField',    'Title of archived document',               'Max: 500, Required'],
    ['content',      'TextField',    'Document content at time of archive',      'Required'],
    ['status',       'CharField',    'Document status at time of archive',       'Max: 50, Required'],
    ['version',      'IntegerField', 'Version number of archived snapshot',      'Required'],
    ['reference_no', 'CharField',    'Document reference number',                'Max: 100, Optional'],
    ['archived_at',  'DateTimeField','Date and time archived',                   'Auto-generated'],
])

# ── documents_document ────────────────────────────────────────────────────────
add_table(doc, 'Table: documents_document', HEADERS, [
    ['id',                   'AutoField',    'Primary key',                       'Auto-generated'],
    ['title',                'TextField',    'Full title of the document',        'Required'],
    ['reference_no',         'CharField',    'Legislative reference number',      'Max: 100, Optional'],
    ['author_id',            'ForeignKey (accounts_user)', 'User who authored the document', 'Cascade delete, Required'],
    ['content',              'TextField',    'Full text content of the document', 'Required'],
    ['doc_type',             'CharField',    'Type of legislative document',      'Max: 15, Choices: ORDINANCE, RESOLUTION, Default: ORDINANCE'],
    ['status',               'CharField',    'Current legislative stage',         'Max: 20, Choices: GHOST, DRAFT, FILED, FIRST_READING, REFERRED, COMMITTEE, SECOND_READING, THIRD_READING, APPROVED, Default: DRAFT'],
    ['referred_committee_id','ForeignKey (committees_committee)', 'Committee the document is referred to', 'Set null on delete, Optional'],
    ['current_version',      'IntegerField', 'Current version number',            'Default: 1'],
    ['hearing_date',         'DateField',    'Scheduled committee hearing date',  'Optional'],
    ['session_included',     'BooleanField', 'Included in session agenda flag',   'Default: False'],
    ['public_participation', 'BooleanField', 'Open for public comments flag',     'Default: False'],
    ['amended_content',      'TextField',    'Amended version of the content',    'Optional'],
    ['amendment_status',     'CharField',    'Status of amendment process',       'Max: 20, Choices: IN_PROGRESS, FINALIZED, NO_AMENDMENTS, Default: IN_PROGRESS, Optional'],
    ['source_barangay_doc_id','ForeignKey (barangay_barangayfiles)', 'Originating barangay submission', 'Set null on delete, Optional'],
    ['embedding',            'VectorField',  'Semantic embedding for RAG search', 'Dimensions: 384, Optional'],
    ['created_at',           'DateTimeField','Date document was created',         'Auto-generated'],
    ['updated_at',           'DateTimeField','Date document was last modified',   'Auto-updated'],
])

# ── documents_amendmentnote ───────────────────────────────────────────────────
add_table(doc, 'Table: documents_amendmentnote', HEADERS, [
    ['id',         'AutoField',    'Primary key',                    'Auto-generated'],
    ['doc_id',     'ForeignKey (documents_document)', 'Document being annotated', 'Cascade delete, Required'],
    ['author_id',  'ForeignKey (accounts_user)', 'User who wrote the note', 'Set null on delete, Optional'],
    ['note',       'TextField',    'Amendment note content',         'Required'],
    ['created_at', 'DateTimeField','Date note was created',          'Auto-generated'],
])

# ── documents_legacydocument ──────────────────────────────────────────────────
add_table(doc, 'Table: documents_legacydocument', HEADERS, [
    ['id',            'AutoField',    'Primary key',                         'Auto-generated'],
    ['title',         'CharField',    'Title of the legacy document',        'Max: 1000, Required'],
    ['reference_no',  'CharField',    'Legislative reference number',        'Max: 100, Optional'],
    ['doc_type',      'CharField',    'Type of legislative document',        'Max: 20, Choices: ORDINANCE, RESOLUTION, Required'],
    ['year',          'IntegerField', 'Year the document was enacted',       'Optional'],
    ['pdf_file',      'FileField',    'Uploaded PDF file',                   'Upload to: legacy_documents/YYYY/, Max: 500, Required'],
    ['extracted_text','TextField',    'OCR-extracted text from PDF',         'Optional'],
    ['embedding',     'VectorField',  'Semantic embedding for RAG search',   'Dimensions: 384, Optional'],
    ['ocr_processed', 'BooleanField', 'OCR processing completed flag',       'Default: False'],
    ['uploaded_at',   'DateTimeField','Date document was uploaded',          'Auto-generated'],
])

# ── document_chunk ────────────────────────────────────────────────────────────
add_table(doc, 'Table: document_chunk', HEADERS, [
    ['id',                  'AutoField',    'Primary key',                        'Auto-generated'],
    ['document_id',         'ForeignKey (documents_document)', 'Parent active document', 'Cascade delete, Optional'],
    ['legacy_document_id',  'ForeignKey (documents_legacydocument)', 'Parent legacy document', 'Cascade delete, Optional'],
    ['chunk_type',          'CharField',    'Type/category of the chunk',         'Max: 50, Required'],
    ['chunk_text',          'TextField',    'Text content of the chunk',          'Required'],
    ['embedding',           'VectorField',  'Semantic embedding for RAG search',  'Dimensions: 384, Required'],
    ['chunk_index',         'IntegerField', 'Position index of chunk in document','Required'],
])

# ── gazette_publiccomment (unmanaged) ─────────────────────────────────────────
add_table(doc, 'Table: gazette_publiccomment  (unmanaged — read-only mirror)', HEADERS, [
    ['id',         'AutoField',    'Primary key',                       'Auto-generated'],
    ['document_id','ForeignKey (documents_document)', 'Document being commented on', 'Required'],
    ['name',       'CharField',    "Commenter's name",                  'Max: 200, Required'],
    ['barangay',   'CharField',    "Commenter's barangay",              'Max: 100, Optional'],
    ['comment',    'TextField',    'Comment content',                   'Required'],
    ['is_approved','BooleanField', 'Moderation approval flag',          'Default: False'],
    ['created_at', 'DateTimeField','Date comment was submitted',        'Auto-generated'],
    ['ip_address', 'GenericIPAddressField', 'IP address of commenter',  'Optional'],
    ['tag',        'CharField',    'Comment type tag',                  'Max: 10, Default: comment'],
    ['replyTo_id', 'ForeignKey (self)', 'Parent comment for replies',   'Optional'],
])

# ── committee_level_committeereport ───────────────────────────────────────────
add_table(doc, 'Table: committee_level_committeereport', HEADERS, [
    ['id',         'AutoField',    'Primary key',                                    'Auto-generated'],
    ['draft_id',   'ForeignKey (documents_document)', 'Document under committee review', 'OneToOne, Cascade delete'],
    ['content',    'TextField',    'Accumulated committee report content (Quill HTML)', 'Default: empty'],
    ['status',     'CharField',    'Committee review outcome',                       'Max: 20, Choices: PENDING, APPROVED, FAILED, Default: PENDING'],
    ['final_pdf',  'FileField',    'Final generated committee report PDF',           'Upload to: final_reports/, Optional'],
    ['created_at', 'DateTimeField','Date report was created',                        'Auto-generated'],
    ['updated_at', 'DateTimeField','Date report was last updated',                   'Auto-updated'],
])

# ── committee_level_hearinglog ────────────────────────────────────────────────
add_table(doc, 'Table: committee_level_hearinglog', HEADERS, [
    ['id',                 'AutoField',    'Primary key',                    'Auto-generated'],
    ['report_id',          'ForeignKey (committee_level_committeereport)', 'Parent committee report', 'Cascade delete, Required'],
    ['hearing_date',       'DateField',    'Date of the hearing',            'Required'],
    ['attendance_notes',   'TextField',    'Notes on attendance',            'Optional'],
    ['outcome',            'CharField',    'Result of the hearing',          'Max: 20, Choices: PENDING, APPROVED, RESET, FAILED, Default: PENDING'],
    ['version_discussed',  'PositiveIntegerField', 'Document version reviewed at hearing', 'Auto-incremented per report, Default: 1'],
    ['created_at',         'DateTimeField','Date log entry was created',     'Auto-generated'],
])

# ── secretariat_session ───────────────────────────────────────────────────────
add_table(doc, 'Table: secretariat_session', HEADERS, [
    ['id',                    'AutoField',    'Primary key',                        'Auto-generated'],
    ['session_number',        'CharField',    'Session number for the term',        'Max: 20, Optional'],
    ['council_number',        'CharField',    'Council term number',                'Max: 10, Optional'],
    ['session_time',          'CharField',    'Scheduled time of session',          'Max: 10, Optional'],
    ['previous_session_date', 'DateField',    'Date of the prior session',          'Optional'],
    ['invocation_by',         'CharField',    'Person delivering the invocation',   'Max: 50, Optional'],
    ['session_date',          'DateField',    'Scheduled date of the session',      'Optional'],
    ['date_started',          'DateTimeField','Actual session start timestamp',     'Auto-generated'],
    ['agenda_finalized_date', 'DateTimeField','Date agenda was finalized',          'Auto-generated'],
    ['order_of_business',     'FileField',    'Generated Order of Business PDF',    'Upload to: session/YYYY/, Optional'],
    ['participants',          'TextField',    'List of session participants',       'Required'],
    ['minutes',               'FileField',    'Uploaded session minutes',           'Upload to: session/minutes/YYYY/, Optional'],
])

# ── audit_auditlog ────────────────────────────────────────────────────────────
add_table(doc, 'Table: audit_auditlog', HEADERS, [
    ['id',         'AutoField',    'Primary key',                          'Auto-generated'],
    ['user_id',    'ForeignKey (accounts_user)', 'User who performed the action', 'Set null on delete, Optional'],
    ['action',     'CharField',    'Type of action performed',             'Max: 20, Choices: CREATE, UPDATE, DELETE, LOGIN, LOGOUT, FILE, MOVE'],
    ['severity',   'CharField',    'Severity level of the action',         'Max: 10, Choices: NORMAL, HIGH, Default: NORMAL'],
    ['target',     'CharField',    'Name/label of the affected record',    'Max: 500, Optional'],
    ['detail',     'TextField',    'Additional context or description',    'Optional'],
    ['ip_address', 'GenericIPAddressField', 'IP address of the actor',     'Optional'],
    ['timestamp',  'DateTimeField','Date and time of the action',          'Auto-generated'],
])

# ── systemadmin_systemsetting ─────────────────────────────────────────────────
add_table(doc, 'Table: systemadmin_systemsetting', HEADERS, [
    ['id',                    'AutoField',    'Primary key',                        'Auto-generated'],
    ['republic_name',         'CharField',    'Name of the republic',               'Max: 255, Default: Republic of the Philippines'],
    ['city_name',             'CharField',    'Name of the city',                   'Max: 255, Default: City of San Juan, Metro Manila'],
    ['office_name',           'CharField',    'Name of the legislative office',     'Max: 255, Default: Office of the Sangguniang Panlungsod'],
    ['system_logo',           'ImageField',   'System branding logo',               'Upload to: branding/, Optional'],
    ['current_council_number','IntegerField', 'Active council term number',         'Default: 8'],
    ['default_venue',         'TextField',    'Default session venue description',  'Default: Session Hall, Room 214'],
    ['maintenance_mode',      'BooleanField', 'System maintenance mode flag',       'Default: False'],
])

# ── Save ──────────────────────────────────────────────────────────────────────
output_path = r'C:\Users\ACER\Desktop\Data_Dictionary_LDMS.docx'
doc.save(output_path)
print(f"Saved: {output_path}")
